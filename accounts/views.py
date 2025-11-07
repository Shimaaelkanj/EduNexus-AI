from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status, permissions
from rest_framework_simplejwt.tokens import RefreshToken
from .serializers import RegisterSerializer, LoginSerializer, LessonSerializer
from .models import Lesson
from rest_framework.parsers import MultiPartParser
import os
from django.core.files.storage import default_storage
from django.conf import settings
from fpdf import FPDF
from django.utils.crypto import get_random_string
from docx import Document
# from pptx import Presentation
# from pptx.util import Inches
from accounts.models import User, Lesson
from rest_framework.permissions import IsAuthenticated
from rest_framework_simplejwt.tokens import AccessToken, TokenError

def validate_jwt_token(token: str):
    """
    Validate and decode a JWT token using SimpleJWT.

    Returns:
        dict: payload data if token is valid
    Raises:
        Exception: if token is invalid or expired
    """
    try:
        access_token = AccessToken(token)
        payload = access_token.payload  # Decoded claims
        return payload
    except TokenError as e:
        raise Exception(f"❌ Invalid or expired token: {str(e)}")


class RegisterView(APIView):
    def post(self, request):
        serializer = RegisterSerializer(data=request.data)
        if serializer.is_valid():
            user = serializer.save(is_active=True)
            return Response(RegisterSerializer(user).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

class LoginView(APIView):
    def post(self, request):
        serializer = LoginSerializer(data=request.data)
        if serializer.is_valid():
            user = serializer.validated_data["user"]
            refresh = RefreshToken.for_user(user)
            return Response({
                "refresh": str(refresh),
                "access": str(refresh.access_token)
            }, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

class ProfileView(APIView):
    def get(self, request):
        return Response({"message": "Profile route working!"})


class LessonView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        lessons = Lesson.objects.all()
        serializer = LessonSerializer(lessons, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request):
        serializer = LessonSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)



UPLOAD_DIR = os.path.join(settings.BASE_DIR, "uploads")
EXPORT_DIR = os.path.join(settings.BASE_DIR, "exports")
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)


# class UploadView(APIView):
#     parser_classes = [MultiPartParser]

#     def post(self, request):
#         file_obj = request.FILES.get('file')
#         if not file_obj:
#             return Response({"error": "No file provided"}, status=status.HTTP_400_BAD_REQUEST)

#         file_path = os.path.join(UPLOAD_DIR, file_obj.name)
#         with open(file_path, 'w', encoding='utf-8') as f:
#             raw_text = file_obj.read().decode('utf-8')
#             f.write(raw_text)

#         return Response({"message": "File uploaded successfully", "filename": file_obj.name}, status=status.HTTP_201_CREATED)


# class AnalyzeCVView(APIView):
#     parser_classes = [MultiPartParser]

#     def post(self, request):
#         file_obj = request.FILES.get('file')
#         if not file_obj:
#             return Response({"detail": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

#         content = file_obj.read().decode('utf-8')
#         lines = content.split("\n")
#         skills = [line for line in lines if any(skill in line for skill in ["Python", "Java", "SQL"])]
#         experience_years = sum([int(s) for s in lines if s.isdigit()])

        
#         cv_result = {"filename": file_obj.name, "skills_found": skills, "experience_years": experience_years}

#         return Response(cv_result)


# class RoadmapView(APIView):
#     def get(self, request, filename):
#         file_path = os.path.join(UPLOAD_DIR, filename)
#         if not os.path.exists(file_path):
#             return Response({"detail": "File not found"}, status=status.HTTP_404_NOT_FOUND)

#         with open(file_path, 'r', encoding='utf-8') as f:
#             content = f.read()
#         lines = content.split("\n")
#         skills_found = [line for line in lines if any(skill in line for skill in ["Python", "Java", "SQL"])]
#         experience_years = sum([int(s) for s in lines if s.isdigit()])

#         roadmap = []
#         if "Python" not in skills_found:
#             roadmap.append("Learn Python basics → advanced Python → projects")
#         if "SQL" not in skills_found:
#             roadmap.append("Learn SQL → practice databases → data projects")
#         if experience_years < 2:
#             roadmap.append("Build 2+ small projects to gain experience")
#         roadmap.append("Apply to internships or entry-level positions relevant to your skills")

#         return Response({"filename": filename, "skills_found": skills_found, "experience_years": experience_years, "roadmap": roadmap})



# class ExportPDFView(APIView):
#     def post(self, request):
#         filename = request.data.get("filename")
#         file_path = os.path.join(UPLOAD_DIR, filename)
#         if not filename or not os.path.exists(file_path):
#             return Response({"detail": "File not found"}, status=status.HTTP_404_NOT_FOUND)

#         with open(file_path, 'r', encoding='utf-8') as f:
#             content = f.read()

#         pdf_path = os.path.join(EXPORT_DIR, f"{filename}.pdf")
#         pdf = FPDF()
#         pdf.add_page()
#         pdf.set_font("Arial", size=12)
#         for line in content.split("\n"):
#             pdf.multi_cell(0, 8, line)
#         pdf.output(pdf_path)
#         return Response({"pdf_path": pdf_path})


# class ExportDOCXView(APIView):
#     def post(self, request):
#         filename = request.data.get("filename")
#         file_path = os.path.join(UPLOAD_DIR, filename)
#         if not filename or not os.path.exists(file_path):
#             return Response({"detail": "File not found"}, status=status.HTTP_404_NOT_FOUND)

#         with open(file_path, 'r', encoding='utf-8') as f:
#             content = f.read()

#         docx_path = os.path.join(EXPORT_DIR, f"{filename}.docx")
#         doc = Document()
#         for line in content.split("\n"):
#             doc.add_paragraph(line)
#         doc.save(docx_path)
#         return Response({"docx_path": docx_path})


# class ExportPPTXView(APIView):
#     def post(self, request):
#         filename = request.data.get("filename")
#         file_path = os.path.join(UPLOAD_DIR, filename)
#         if not filename or not os.path.exists(file_path):
#             return Response({"detail": "File not found"}, status=status.HTTP_404_NOT_FOUND)

#         with open(file_path, 'r', encoding='utf-8') as f:
#             content = f.read()

#         pptx_path = os.path.join(EXPORT_DIR, f"{filename}.pptx")
#         prs = Presentation()
#         slide_layout = prs.slide_layouts[1]
#         lines = content.split("\n")
#         for i in range(0, len(lines), 5):
#             slide = prs.slides.add_slide(slide_layout)
#             slide.shapes.title.text = f"{filename} - Slide {i//5 + 1}"
#             slide.placeholders[1].text = "\n".join(lines[i:i+5])
#         prs.save(pptx_path)
#         return Response({"pptx_path": pptx_path})
    
